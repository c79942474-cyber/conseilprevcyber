# -*- coding: utf-8 -*-
"""Remplir les formulaires officiels DANS LEUR PROPRE FICHIER, sans jamais
signer ni déclarer.

CE QUE CE MODULE FAIT, ET POURQUOI IL NE REDESSINE RIEN. `ao_dc.remplir()`
produit un report tracé : chaque rubrique, sa valeur, son origine. Ce report se
lit à côté du formulaire ; il ne se dépose pas à sa place. Ici, on va plus
loin : on ouvre LE FORMULAIRE OFFICIEL — le fichier du ministère, dans sa
version, avec sa mise en page et sa date de mise à jour — et on écrit les
valeurs dedans.

LA DIFFÉRENCE AVEC UN FAC-SIMILÉ EST TOUT L'ENJEU. Un formulaire redessiné par
un programme serait refusé, ou pire, accepté et faux. Un formulaire OFFICIEL
rempli reste le formulaire officiel : rien de sa mise en page, de ses mentions
légales ou de sa numérotation de cadres n'est réécrit. On ne touche qu'aux
emplacements que le modèle laisse vides.

TROIS INVARIANTS, ET CE SONT DES INVARIANTS, PAS DES INTENTIONS :

  1. ON N'ÉCRIT QUE DANS UN EMPLACEMENT VIDE. Un paragraphe qui porte du texte
     n'est jamais modifié. Cela rend STRUCTURELLEMENT impossible d'écraser une
     déclaration sur l'honneur, une mention légale ou un intitulé de cadre —
     là où une liste noire de zones à éviter serait une énumération, donc
     oubliable.

  2. LES ZONES DE DÉCLARATION ET DE SIGNATURE SONT HORS D'ATTEINTE. Même vides,
     leurs emplacements sont exclus : une ancre qui glisserait ne peut pas y
     déposer une valeur. C'est la seconde barrière, et elle est là parce que la
     première protège du texte, pas du blanc.

  3. SEULES LES VALEURS AU STATUT « rempli » SONT ÉCRITES. Une déclaration
     ressort toujours au statut `a_declarer` sans valeur : elle ne peut donc
     pas être écrite, même si une ancre la désignait. Troisième barrière,
     posée en amont des deux autres.

L'EMPREINTE DU MODÈLE EST VÉRIFIÉE AVANT D'ÉCRIRE. Les ancres sont des phrases
du formulaire ; un formulaire mis à jour peut les déplacer, les reformuler ou
les supprimer — et le remplissage écrirait alors à côté, sans rien signaler.
Le modèle est donc épinglé par son empreinte : un fichier qui a changé fait
REFUSER le remplissage, avec la raison. Mettre à jour le modèle oblige à
revérifier les ancres, ce qui est exactement le geste qu'on veut imposer.

CE MODULE NE SIGNE RIEN, ET LE DOCUMENT LE DIT. Le fichier produit porte en
tête une ligne qui le nomme pour ce qu'il est : un projet rempli
automatiquement, non signé et non vérifié. Un document qui circule sans dire
cela finirait par être déposé tel quel.
"""
import hashlib
import io
import os
import re

VERSION = "2026-08-a"

ICI = os.path.dirname(os.path.abspath(__file__))
DOSSIER_MODELES = os.path.join(ICI, "modeles")

# ── LA LIGNE QUE PORTE CHAQUE DOCUMENT PRODUIT ────────────────────────────
# Elle est ajoutée EN TÊTE, jamais insérée dans un cadre : le formulaire n'est
# pas altéré, il est précédé. Elle se retire d'un coup de touche avant
# signature — et si on l'oublie, on dépose un formulaire qui dit qu'il est un
# projet, ce qui est gênant mais honnête. L'inverse ne l'est pas.
BANDEAU = ("PROJET — rempli automatiquement à partir de votre fiche et des "
           "pièces de consultation déposées. NON SIGNÉ, NON VÉRIFIÉ. Les "
           "déclarations sur l'honneur et les blocs de signature sont restés "
           "vides : personne ici ne déclare et ne signe à votre place. "
           "Relisez, complétez, puis supprimez cette ligne.")

# ── LES MODÈLES, ÉPINGLÉS PAR LEUR EMPREINTE ──────────────────────────────
# `maj` est la date que le formulaire porte LUI-MÊME, en dernière page. Elle
# n'est pas décorative : c'est ce qui permet de dire à l'utilisateur quelle
# version il remplit, et de constater qu'elle a vieilli.
MODELES = {
    "dc4": {
        "fichier": "DC4.docx",
        "nom": "DC4 — Déclaration de sous-traitance",
        "piece": "dc4",
        "maj": "12/10/2023",
        "source": "Direction des affaires juridiques, ministère chargé de "
                  "l'économie — formulaire type, document facultatif.",
        "empreinte": ("b39e32506f672bfc4a37c7121a47e1b7b61eb39b3b73272e"
                      "cc008945a2dd860a"),
    },
}

# ── OÙ CHAQUE VALEUR SE POSE ──────────────────────────────────────────────
# L'ancre est une phrase DU FORMULAIRE, recopiée telle qu'il l'écrit. La valeur
# va dans le premier emplacement vide qui suit, sans jamais franchir l'intitulé
# du cadre suivant : une ancre qui ne trouve rien avant le cadre d'après ne
# place RIEN et le dit. Écrire « quelque part par là » serait pire que ne rien
# écrire, parce que la case aurait l'air remplie.
ANCRES = {
    # LES CADRES D ET E PORTENT LES MÊMES INTITULÉS, MOT POUR MOT : « Adresse
    # électronique : », « Numéro SIRET… », « Nom commercial et dénomination
    # sociale… ». L'un identifie le TITULAIRE, l'autre le SOUS-TRAITANT. Une
    # ancre sans rang déposerait donc le SIRET du titulaire dans la case du
    # sous-traitant — une faute que personne ne verrait, puisque la case serait
    # pleine et plausible. `occurrence` est là pour cela, et une règle la
    # mesure sur le document produit.
    "dc4": [
        {"rubrique": "acheteur", "ancre": "Désignation de l’acheteur :"},
        {"rubrique": "objet_marche", "ancre": "B - Objet du marché public"},

        # ── cadre D — le titulaire, premier rang ──────────────────────────
        {"rubrique": "titulaire", "occurrence": 1,
         "ancre": "Nom commercial et dénomination sociale de l’unité ou de "
                  "l’établissement qui exécutera la prestation :"},
        {"rubrique": "titulaire_adresse", "occurrence": 1,
         "ancre": "Adresses postale et du siège social (si elle est "
                  "différente de l’adresse postale) :"},
        {"rubrique": "titulaire_courriel", "occurrence": 1,
         "ancre": "Adresse électronique :"},
        {"rubrique": "titulaire_siret", "occurrence": 1,
         "ancre": "Numéro SIRET, à défaut, un numéro d’identification "
                  "européen"},
        {"rubrique": "titulaire_forme", "occurrence": 1,
         "ancre": "Forme juridique du soumissionnaire individuel, du "
                  "titulaire ou du membre du groupement"},
        {"rubrique": "mandataire", "occurrence": 1,
         "ancre": "En cas de groupement momentané d’entreprises, "
                  "identification et coordonnées du mandataire"},

        # ── cadre E — le sous-traitant, second rang ───────────────────────
        {"rubrique": "sous_traitant", "occurrence": 2,
         "ancre": "Nom commercial et dénomination sociale de l’unité ou de "
                  "l’établissement qui exécutera la prestation :"},
        {"rubrique": "sous_traitant_pouvoir", "occurrence": 1,
         "ancre": "Personne(s) physique(s) ayant le pouvoir d’engager le "
                  "sous-traitant :"},

        # ── cadres F à J ──────────────────────────────────────────────────
        {"rubrique": "prestations",
         "ancre": "Nature des prestations sous-traitées :"},
        {"rubrique": "montant",
         "ancre": "Montant des prestations sous-traitées :"},
        {"rubrique": "variation_prix",
         "ancre": "Modalités de variation des prix :"},
        {"rubrique": "compte",
         "ancre": "Nom de l’établissement bancaire :"},
        {"rubrique": "duree_sous_traitance",
         "ancre": "La durée du contrat de sous-traitance en nombre de mois "
                  "est de :"},
        {"rubrique": "capacites",
         "ancre": "J1 - Récapitulatif des informations et renseignements"},
    ],
}

# ── CE QUI EST HORS D'ATTEINTE, MÊME VIDE ─────────────────────────────────
# Chaque entrée ouvre une zone qui court jusqu'à l'intitulé de cadre suivant —
# ou jusqu'à la fin du document quand `jusqu_a_la_fin` est vrai. Les
# déclarations sur l'honneur et les signatures ne se remplissent pas : elles
# s'assument.
ZONES_INTERDITES = {
    "dc4": [
        {"ancre": "K1 - Le sous-traitant déclare sur l’honneur",
         "jusqu_a_la_fin": False},
        {"ancre": "M - Acceptation et agrément des conditions de paiement",
         "jusqu_a_la_fin": True},
    ],
}

_CADRE = re.compile(r"^[A-N][0-9]?\s*[-–]\s+\S")
_VIDE = re.compile(r"^[\s.…_· ]*$")


def _net(s):
    """Un texte comparable : espaces normalisés, apostrophes unifiées.

    LES APOSTROPHES SONT LE PIÈGE. Le formulaire écrit « l’acheteur » avec une
    apostrophe typographique ; une ancre saisie au clavier porte souvent
    l'apostrophe droite. Deux chaînes qui se ressemblent à l'œil ne se
    comparent pas — et l'ancre resterait introuvable sans que rien n'explique
    pourquoi.
    """
    return " ".join((s or "").replace("’", "'").replace("ʼ", "'")
                    .split())


def _emplacement_libre(texte):
    """Un emplacement à remplir : vide, ou fait de pointillés."""
    return bool(_VIDE.match((texte or "").replace("’", "'")))


def empreinte(chemin):
    """L'empreinte du fichier — celle qui dit si le modèle a bougé."""
    with io.open(chemin, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def chemin_modele(cle):
    return os.path.join(DOSSIER_MODELES, MODELES[cle]["fichier"])


def modeles_disponibles():
    """Les modèles présents ET intacts, et ceux qui manquent — dits à part.

    UN MODÈLE ABSENT N'EST PAS UN MODÈLE ALTÉRÉ. Les confondre enverrait
    chercher une corruption là où il n'y a qu'un fichier à déposer.
    """
    prets, manquants, alteres = [], [], []
    for cle, m in sorted(MODELES.items()):
        p = chemin_modele(cle)
        if not os.path.exists(p):
            manquants.append(cle)
        elif empreinte(p) != m["empreinte"]:
            alteres.append(cle)
        else:
            prets.append(cle)
    return {"prets": prets, "manquants": manquants, "alteres": alteres}


def _zones_interdites(paras, cle):
    """Les indices de paragraphes où l'on n'écrira jamais."""
    interdits = set()
    for z in ZONES_INTERDITES.get(cle, []):
        cible = _net(z["ancre"])
        for i, t in enumerate(paras):
            if not _net(t).startswith(cible):
                continue
            j = i + 1
            while j < len(paras):
                if not z["jusqu_a_la_fin"] and _CADRE.match(_net(paras[j])):
                    break
                interdits.add(j)
                j += 1
            interdits.add(i)
    return interdits


def _cible(paras, interdits, pris, ancre, occurrence):
    """L'indice de l'emplacement à remplir, ou None — jamais un « à peu près ».

    ON NE FRANCHIT PAS L'INTITULÉ DU CADRE SUIVANT. Une ancre dont
    l'emplacement a disparu chercherait sinon jusqu'au bout du document et
    déposerait sa valeur dans un cadre qui ne l'attend pas.

    INTERDIT ET DÉJÀ PRIS NE SE TRAITENT PAS PAREIL, et les confondre serait
    une faute dans les deux sens. Un emplacement INTERDIT fait renoncer :
    glisser plus loin reviendrait à contourner la barrière qu'on vient de
    poser. Un emplacement DÉJÀ PRIS fait simplement passer au suivant — un
    cadre offre souvent plusieurs lignes vides d'affilée, et refuser la
    seconde parce que la première est occupée laisserait une case vide sans
    raison.
    """
    cible = _net(ancre)
    vues = 0
    for i in range(len(paras)):
        if not _net(paras[i]).startswith(cible):
            continue
        vues += 1
        if vues < occurrence:
            continue
        for j in range(i + 1, len(paras)):
            if _CADRE.match(_net(paras[j])):
                return None
            if j in interdits:
                return None
            if j in pris:
                continue
            if _emplacement_libre(paras[j]):
                return j
        return None
    return None


def remplir_document(cle, valeurs, bandeau=BANDEAU):
    """Le formulaire officiel, rempli de ce qui est déjà connu. Rend
    (octets, rapport).

    `valeurs` est un dictionnaire {clé de rubrique: valeur} — celles que
    `ao_dc.remplir()` a portées au statut « rempli ». Rien d'autre n'entre.

    LE RAPPORT DIT CE QUI A ÉTÉ PLACÉ ET CE QUI NE L'A PAS ÉTÉ. Un formulaire
    rendu sans dire ce qui manque se lit comme un formulaire complet.
    """
    from docx import Document                                   # noqa: PLC0415

    m = MODELES[cle]
    p = chemin_modele(cle)
    if not os.path.exists(p):
        return None, {"ok": False, "motif": "modele_absent", "modele": cle,
                      "places": [], "non_places": [], "ignores": []}
    reelle = empreinte(p)
    if reelle != m["empreinte"]:
        # ON REFUSE PLUTÔT QUE DE REMPLIR À CÔTÉ. Les ancres sont des phrases
        # de CE fichier-ci ; un modèle mis à jour les déplace sans prévenir.
        return None, {"ok": False, "motif": "modele_altere", "modele": cle,
                      "attendue": m["empreinte"], "trouvee": reelle,
                      "places": [], "non_places": [], "ignores": []}

    doc = Document(p)
    paras = [pa.text for pa in doc.paragraphs]
    interdits = _zones_interdites(paras, cle)
    places, non_places, ignores = [], [], []
    pris = set()

    for a in ANCRES.get(cle, []):
        v = str(valeurs.get(a["rubrique"]) or "").strip()
        if not v:
            ignores.append({"rubrique": a["rubrique"], "motif": "sans_valeur"})
            continue
        i = _cible(paras, interdits, pris, a["ancre"], a.get("occurrence", 1))
        if i is None:
            non_places.append({"rubrique": a["rubrique"], "ancre": a["ancre"],
                               "motif": "emplacement_introuvable"})
            continue
        cible = doc.paragraphs[i]
        for r in list(cible.runs):
            r._element.getparent().remove(r._element)
        cible.add_run(v)
        paras[i] = v
        pris.add(i)
        places.append({"rubrique": a["rubrique"], "ancre": a["ancre"],
                       "valeur": v, "paragraphe": i})

    if bandeau:
        # EN TÊTE, ET DANS UN PARAGRAPHE À LUI. Insérer le bandeau dans un
        # cadre du formulaire modifierait le formulaire ; le poser devant ne le
        # modifie pas.
        premier = doc.paragraphs[0]
        neuf = premier.insert_paragraph_before(bandeau)
        for r in neuf.runs:
            r.bold = True

    flux = io.BytesIO()
    doc.save(flux)
    return flux.getvalue(), {
        "ok": True, "motif": "ok", "modele": cle, "nom": m["nom"],
        "maj": m["maj"], "source": m["source"], "empreinte": reelle,
        "places": places, "non_places": non_places, "ignores": ignores,
        "bandeau": bool(bandeau),
    }


def valeurs_pour(remplissage, piece):
    """Ce qu'on peut écrire, pris dans le report de `ao_dc.remplir()`.

    SEUL CE QUI EST « rempli » PASSE. Une déclaration est toujours
    `a_declarer` et sans valeur ; une case à saisir est `a_saisir`. Ni l'une ni
    l'autre n'entre dans le document — la première parce que personne ne
    déclare à votre place, la seconde parce qu'il n'y a rien à écrire.
    """
    for p in (remplissage or {}).get("pieces", []):
        if p["cle"] != piece:
            continue
        return {l["cle"]: l["valeur"] for l in p["rubriques"]
                if l["statut"] == "rempli" and l["valeur"]}
    return {}
