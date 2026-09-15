# -*- coding: utf-8 -*-
"""LE CADRE E DU DC1 SE REMPLIT TOUT SEUL, DEPUIS LA NOTE DE RÉPARTITION.

CE QUI ÉTAIT EN CAUSE. Le cadre E du DC1 — « Identification des membres du
groupement et répartition des prestations » — est une GRILLE, pas une suite de
paragraphes. Le remplissage du module marche par ancres textuelles dans le
corps du document : il n'avait aucun moyen d'y écrire, et le cadre E n'avait
JAMAIS été rempli par quoi que ce soit. Un candidat en groupement recopiait ses
quatre cotraitants à la main, dans un tableau, à chaque consultation.

CE QUE CES RÈGLES TIENNENT, ET QU'AUCUNE LECTURE DE SOURCE NE DONNE :

  · LA LISTE VIENT DE LA NOTE, ET DE NULLE PART AILLEURS. Elle n'est ni
    devinée, ni complétée, ni réordonnée. Une note sans en-tête reconnaissable
    ne rend RIEN.

  · LE CÔTÉ DÉCIDE. La même note déposée du côté ACHETEUR ne verse pas : ce
    serait le groupement d'un concurrent qui entrerait dans notre DC1.

  · LE DOCUMENT PRODUIT EST MESURÉ EN L'OUVRANT. Pas le rapport, pas le
    Markdown : le .docx, rouvert, cellule par cellule. C'est la seule mesure
    qui distingue « la fonction a rendu 5 » de « le formulaire porte 5 noms ».

  · CE QUI N'EST PAS ÉCRIT NE L'EST PAS EN SILENCE. Le n° de lot et les
    coordonnées de chaque membre restent vides, parce que la note ne les donne
    pas, et le rapport les NOMME. Une case à moitié remplie qui a l'air
    complète est pire qu'une case vide.

  · LES RÉSERVES DE L'AUTEUR SURVIVENT À LA RECOPIE. « [À confirmer] » part
    dans le formulaire tel quel, et il est COMPTÉ : un formulaire officiel où
    une réserve de travail passe pour un engagement pris est un faux poli.
"""
import io
import json
import os
import re
import sys
import zipfile

ICI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ICI)

import ao_dc                                                       # noqa: E402
import ao_formulaires                                              # noqa: E402
from conftest import ORIGINE                                       # noqa: E402
from test_ao_formulaires import _js_source                         # noqa: E402

import pytest                                                      # noqa: E402

FICHE = {"raison_sociale": "CONSEILPREV", "siret": "73282932000074"}

# ── LA NOTE D'ESSAI, ÉCRITE ICI ─────────────────────────────────────────────
# ELLE EST APLATIE COMME `rag_store.extract_text` APLATIT UN DOCX : une ligne
# par rangée, cellules jointes par « | ». Écrire un vrai .docx d'essai aurait
# mesuré python-docx ; ce qu'on veut mesurer, c'est CE QUE LE MODULE FAIT du
# texte qu'il reçoit, et ce texte-là a exactement cette forme.
NOTE = u"""Note de répartition des compétences — groupement conjoint

Le présent document fixe la répartition des prestations entre les membres.

Membre | Rôle | Compétences et prestations
EPOC Ingénierie | Mandataire — architecte | Conception architecturale
EPR Ingénierie | Cotraitant — BET fluides | [À confirmer] Courants forts
DEAR Concept | Cotraitant — BET structure | [À confirmer] Structure
CONSEILPREV | Cotraitant — cybersécurité | Sécurité des systèmes industriels

La convention de groupement précise la solidarité du mandataire.
"""

NOMS_ATTENDUS = ["EPOC Ingénierie", "EPR Ingénierie", "DEAR Concept",
                 "CONSEILPREV"]
AVEC_RESERVE = ["EPR Ingénierie", "DEAR Concept"]

FICHIER_NOTE = "Note-repartition-competences-groupement.docx"


def _document(docs):
    return [dict(d, extension=".docx") for d in docs]


def _analyse_avec_note(cote="cabinet", nom=FICHIER_NOTE, texte=NOTE):
    return ao_dc.analyser(_document([{"nom": nom, "texte": texte,
                                      "cote": cote}]))


def _produire(membres):
    """Le DC1 rempli, rendu comme un document ROUVERT — pas comme un rapport."""
    from docx import Document
    octets, rap = ao_formulaires.remplir_document(
        "dc1", {}, membres=membres)
    assert rap.get("ok"), rap.get("motif")
    return Document(io.BytesIO(octets)), rap


def _rang_grille(doc):
    """LE RANG de la grille du cadre E parmi les tables, ou None.

    ON RENVOIE UN INDICE, PAS L'OBJET. `doc.tables` fabrique un proxy neuf à
    chaque accès : comparer deux tables par identité rend toujours faux, et une
    règle écrite ainsi passerait sans rien mesurer."""
    for i, t in enumerate(doc.tables):
        tete = " ".join(" ".join(c.text.lower() for c in t.rows[0].cells).split())
        if "nom commercial" in tete and "prestations" in tete:
            return i
    return None


def _grille(doc):
    """La grille du cadre E, retrouvée sur ses intitulés de colonnes."""
    i = _rang_grille(doc)
    return None if i is None else doc.tables[i]


def _modele():
    from docx import Document
    return Document(ao_formulaires.chemin_modele("dc1"))


# ══════════════════════════════════════════════════════════════════════════
#  1. LA NOTE EST RECONNUE — SUR UN NOM DE FICHIER, PAS SUR UNE PHRASE
# ══════════════════════════════════════════════════════════════════════════

def test_la_note_de_repartition_est_rattachee_a_sa_piece():
    """MESURÉ AVANT CORRECTION : `piece_du_cabinet` rendait None sur ce nom.
    `CABINET_MOTIFS` n'avait aucune entrée pour la répartition des
    compétences — seize des vingt-trois pièces en avaient une — si bien qu'une
    note parfaitement nommée s'affichait « son nom ne le rattache à aucune des
    pièces à produire », et que le cadre E restait vide en silence."""
    assert ao_dc.piece_du_cabinet(FICHIER_NOTE) == "repartition_competences"


def test_le_motif_tolere_le_tiret_ET_l_absence_d_article():
    """UN NOM DE FICHIER N'EST PAS UNE PHRASE. Il s'écrit sans espace, sans
    accent et sans article. Un motif qui exige « répartition DES compétences »
    est vrai sur le papier et inerte sur le disque : c'est exactement ce qui a
    fait échouer la première version."""
    for nom in ("repartition-competences.docx",
                "Repartition_des_taches_2026.docx",
                "note repartition des prestations.pdf",
                "REPARTITION-MISSIONS-GROUPEMENT.docx"):
        assert ao_dc.piece_du_cabinet(nom) == "repartition_competences", nom


def test_chaque_piece_du_catalogue_atteignable_a_un_motif():
    """LES SEULES SANS MOTIF SONT LES QUATRE FORMULAIRES DE L'ÉTAT, et leur
    absence est VOULUE : ce module les produit, un exemplaire vierge déposé par
    l'acheteur est relevé par `_formulaire_fourni`, et leur donner un motif du
    côté cabinet ferait passer notre propre production pour une pièce apportée.

    CE QUE CETTE RÈGLE ATTRAPE : une pièce ajoutée au catalogue sans motif —
    donc indéposable, silencieusement."""
    cles = [p["cle"] for p, _d in ao_dc._catalogue()]
    sans = [c for c in cles if c not in ao_dc.CABINET_MOTIFS]
    assert sorted(sans) == sorted(["dc1", "dc2", "dc4", "acte_engagement"]), \
        "pièce(s) du catalogue qu'aucun nom de fichier ne peut atteindre : " \
        + ", ".join(sorted(sans))


# ══════════════════════════════════════════════════════════════════════════
#  2. LA GRILLE SE LIT — ET NE SE DEVINE PAS
# ══════════════════════════════════════════════════════════════════════════

def test_les_membres_sont_lus_dans_l_ordre_de_la_note():
    """L'ORDRE EST UNE INFORMATION : le mandataire vient en premier, et le
    cadre E se lit comme la note. Le trier alphabétiquement ferait passer un
    cotraitant pour le mandataire."""
    lus = ao_dc.membres_du_groupement(NOTE)
    assert [m["nom"] for m in lus] == NOMS_ATTENDUS


def test_les_reserves_de_l_auteur_sont_portees_par_la_ligne():
    """« [À confirmer] » n'est pas un détail de mise en forme : c'est l'auteur
    de la note qui dit que cette répartition n'est pas arrêtée. Deux lecteurs
    de la liste qui le recalculeraient chacun de leur côté en tireraient deux
    comptes différents — la ligne le porte donc elle-même."""
    lus = ao_dc.membres_du_groupement(NOTE)
    assert [m["nom"] for m in lus if m["a_confirmer"]] == AVEC_RESERVE


def test_sans_entete_reconnaissable_rien_n_est_rendu():
    """DEVINER SERAIT LE PIRE. Un planning, une liste de prix, un tableau
    d'honoraires ont trois colonnes eux aussi. Les prendre pour la répartition
    écrirait n'importe quoi dans un formulaire de l'État."""
    planning = u"""Phase | Début | Fin
APS | janvier | mars
APD | mars | juin
"""
    assert ao_dc.membres_du_groupement(planning) == []


def test_un_tableau_a_deux_colonnes_n_est_pas_la_repartition():
    """LA LARGEUR EST LE PREMIER GARDE-FOU. « Membre | Rôle » existe dans
    d'autres notes — un organigramme, une liste de contacts — et le cadre E
    attend trois valeurs par rangée : un tableau à deux colonnes ne peut pas
    les donner.

    CE QUE CETTE RÈGLE NE MESURE PAS : l'exigence des trois INTITULÉS, qui est
    un second garde-fou et qui a sa propre règle juste en dessous. Les deux
    tombaient ensemble sur ce seul cas, et la seconde ne mesurait donc rien."""
    contacts = u"""Membre | Rôle
Paul Durand | Directeur
Marie Leroy | Responsable QSE
"""
    assert ao_dc.membres_du_groupement(contacts) == []


def test_trois_colonnes_ne_suffisent_PAS_il_faut_les_trois_intitules():
    """LE CAS QUI SÉPARE LES DEUX GARDE-FOUS : un tableau de contacts À TROIS
    COLONNES. Il passe la largeur, et seule l'exigence du troisième intitulé —
    compétence, ou prestation — l'arrête.

    CE QU'ON ÉVITE : un DC1 dont le cadre E nomme trois PERSONNES et leur
    numéro de téléphone, là où l'acheteur attend les entreprises cotraitantes
    et ce que chacune exécute."""
    contacts = u"""Membre | Rôle | Téléphone
Paul Durand | Directeur | 01 02 03 04 05
Marie Leroy | Responsable QSE | 01 02 03 04 06
"""
    assert ao_dc.membres_du_groupement(contacts) == []


def test_la_lecture_s_arrete_ou_le_tableau_s_arrete():
    """La prose qui suit la grille n'a pas de « | » : c'est ce qui borne la
    lecture sans avoir à deviner combien de membres le groupement compte. Sans
    cette borne, la phrase sur la convention de groupement deviendrait un
    cinquième cotraitant."""
    lus = ao_dc.membres_du_groupement(NOTE)
    assert len(lus) == 4
    assert all("convention" not in m["nom"].lower() for m in lus)


# ══════════════════════════════════════════════════════════════════════════
#  3. LE CÔTÉ DÉCIDE — ET C'EST CE QUI PROTÈGE LE DC1
# ══════════════════════════════════════════════════════════════════════════

def test_la_note_deposee_du_cote_cabinet_verse_les_membres():
    """LE FIL, PAR LES VRAIES FONCTIONS. Reproduire ici les ingrédients de
    `analyser()` ferait passer la règle pendant que la route rendrait autre
    chose — c'est le défaut qu'une mutation a déjà fait apparaître ailleurs
    dans ce dossier."""
    a = _analyse_avec_note(cote="cabinet")
    assert [m["nom"] for m in a["membres_groupement"]] == NOMS_ATTENDUS
    ligne = a["pieces_cabinet"][0]
    assert ligne["cle"] == "repartition_competences"
    assert len(ligne["membres"]) == 4


def test_la_MEME_note_deposee_du_cote_acheteur_ne_verse_RIEN():
    """LE GROUPEMENT D'UN CONCURRENT NE DOIT PAS ENTRER DANS NOTRE DC1.

    Une note de répartition trouvée dans le dossier de consultation — jointe
    par l'acheteur en exemple, ou déposée du mauvais côté — décrit le
    groupement de quelqu'un d'autre. L'écrire dans notre cadre E produirait un
    formulaire officiel nommant des entreprises qui n'ont rien signé."""
    a = _analyse_avec_note(cote="consultation")
    assert a["membres_groupement"] == []
    assert not any(l.get("membres") for l in a.get("pieces_cabinet") or [])


def test_le_premier_fichier_qui_porte_une_grille_gagne():
    """Deux notes déposées, c'est la PREMIÈRE qui fait foi — la même règle que
    `fournies_cabinet`. Prendre la dernière ferait dépendre le cadre E de
    l'ordre dans lequel les fichiers ont été glissés dans le navigateur."""
    autre = NOTE.replace("EPOC Ingénierie", "ZZZ Ingénierie")
    a = ao_dc.analyser(_document([
        {"nom": FICHIER_NOTE, "texte": NOTE, "cote": "cabinet"},
        {"nom": "Repartition-competences-v2.docx", "texte": autre,
         "cote": "cabinet"}]))
    assert a["membres_groupement"][0]["nom"] == "EPOC Ingénierie"


def test_remplir_porte_les_membres_jusqu_au_DC1_et_a_lui_seul():
    """LA NOTE NE NOURRIT QUE LE CADRE E. Les porter sur les vingt-trois
    pièces ferait croire qu'elle alimente tout le dossier."""
    r = ao_dc.remplir(fiche=FICHE, analyse=_analyse_avec_note(),
                      groupement=True)
    assert [m["nom"] for m in r["membres_groupement"]] == NOMS_ATTENDUS
    porteuses = [p["cle"] for p in r["pieces"] if p.get("membres")]
    assert porteuses == ["dc1"]
    assert ao_formulaires.membres_pour(r) == r["membres_groupement"]


def test_sans_note_deposee_aucun_membre_n_est_invente():
    """LE TÉMOIN NÉGATIF. Sans lui, une règle qui compte 4 membres passerait
    aussi bien sur un module qui en écrit 4 quoi qu'on lui donne."""
    r = ao_dc.remplir(fiche=FICHE, analyse=ao_dc.analyser([]), groupement=True)
    assert r["membres_groupement"] == []
    assert ao_formulaires.membres_pour(r) == []


# ══════════════════════════════════════════════════════════════════════════
#  4. LE DOCUMENT PRODUIT — MESURÉ EN L'OUVRANT
# ══════════════════════════════════════════════════════════════════════════

def test_le_DC1_produit_porte_un_membre_par_ligne_de_grille():
    """ON ROUVRE LE .docx. Le rapport dit ce que la fonction CROIT avoir fait ;
    seule la lecture du fichier dit ce que l'acheteur recevra."""
    doc, rap = _produire(ao_dc.membres_du_groupement(NOTE))
    t = _grille(doc)
    assert t is not None, "la grille du cadre E n'existe plus dans le modèle"
    lignes = [r.cells[1].text.strip() for r in t.rows[1:]]
    assert [x for x in lignes if x] == NOMS_ATTENDUS
    assert rap["cadre_e"]["ecrits"] == 4


def test_les_prestations_de_chaque_membre_suivent_leur_nom():
    """UNE RÉPARTITION DÉCALÉE D'UNE LIGNE EST LE PIRE DES RÉSULTATS : elle se
    découvre à l'exécution, quand un cotraitant refuse une tâche que le DC1 lui
    attribue."""
    membres = ao_dc.membres_du_groupement(NOTE)
    doc, _ = _produire(membres)
    t = _grille(doc)
    for i, m in enumerate(membres):
        rang = t.rows[i + 1]
        assert rang.cells[1].text.strip() == m["nom"]
        assert rang.cells[2].text.strip() == m["prestations"]


def test_les_reserves_partent_TELLES_QUELLES_dans_le_formulaire():
    """LES EFFACER EN RECOPIANT ferait passer pour arrêté ce que l'auteur a
    signalé comme ouvert — dans un document signé par une personne habilitée."""
    doc, rap = _produire(ao_dc.membres_du_groupement(NOTE))
    t = _grille(doc)
    porteuses = [r.cells[1].text.strip() for r in t.rows[1:]
                 if "[À confirmer]" in r.cells[2].text]
    assert porteuses == AVEC_RESERVE
    assert rap["cadre_e"]["a_confirmer"] == AVEC_RESERVE


def test_le_numero_de_lot_reste_VIDE_parce_que_la_note_ne_le_donne_pas():
    """DÉDUIRE « 1 » SERAIT UNE INVENTION. Un marché alloti attribue des lots
    différents aux membres ; écrire un numéro que personne n'a donné engage le
    groupement sur un périmètre qu'il n'a pas choisi. La colonne reste vide, et
    le rapport DIT qu'elle reste à compléter."""
    doc, rap = _produire(ao_dc.membres_du_groupement(NOTE))
    t = _grille(doc)
    assert all(not r.cells[0].text.strip() for r in t.rows[1:])
    assert any("lot" in x.lower() for x in rap["cadre_e"]["a_completer"])


def test_ce_que_le_cadre_demande_encore_est_NOMME():
    """Le cadre E veut aussi l'adresse, le courriel, le téléphone et le SIRET
    de chaque membre. Ni la note ni notre fiche ne les tiennent pour un
    cotraitant. Une case à moitié remplie qui a l'air complète est pire qu'une
    case vide : on ne la relit pas."""
    _doc, rap = _produire(ao_dc.membres_du_groupement(NOTE))
    dit = " ".join(rap["cadre_e"]["a_completer"]).lower()
    for mot in ("adresse", "siret"):
        assert mot in dit, "le rapport ne nomme pas « %s »" % mot


def test_la_grille_grandit_du_nombre_exact_de_membres():
    """LE MODÈLE OFFRE QUATRE RANGÉES LIBRES. Un groupement de six ne doit ni
    perdre deux membres, ni faire grandir la grille de six."""
    modele = _grille(_modele())
    libres = len(modele.rows) - 1
    membres = [{"nom": "M%d" % i, "role": "", "prestations": "p%d" % i}
               for i in range(libres + 2)]
    doc, rap = _produire(membres)
    t = _grille(doc)
    assert rap["cadre_e"]["lignes_ajoutees"] == 2
    assert len(t.rows) - 1 == len(membres)
    assert [r.cells[1].text.strip() for r in t.rows[1:]] == \
        [m["nom"] for m in membres]


def test_une_grille_plus_grande_que_besoin_garde_ses_rangees_vides():
    """MOINS DE MEMBRES QUE DE RANGÉES : les rangées en trop restent vides. Y
    écrire quoi que ce soit — un tiret, un « néant » — ferait compter un
    cotraitant de plus à la lecture."""
    doc, rap = _produire([{"nom": "Seule SARL", "role": "Mandataire",
                           "prestations": "Tout"}])
    t = _grille(doc)
    assert rap["cadre_e"]["lignes_ajoutees"] == 0
    assert [r.cells[1].text.strip() for r in t.rows[1:]] == \
        ["Seule SARL"] + [""] * (len(t.rows) - 2)


def test_le_nombre_de_membres_est_plafonne():
    """UNE NOTE MAL FORMÉE POURRAIT RENDRE DES CENTAINES DE RANGÉES, et le
    document deviendrait illisible — voire impossible à ouvrir. Le plafond est
    déclaré, il n'est pas implicite."""
    doc, rap = _produire([{"nom": "M%d" % i, "role": "", "prestations": ""}
                          for i in range(ao_formulaires.MEMBRES_MAX + 25)])
    assert rap["cadre_e"]["ecrits"] == ao_formulaires.MEMBRES_MAX
    assert len(_grille(doc).rows) - 1 == ao_formulaires.MEMBRES_MAX


# ══════════════════════════════════════════════════════════════════════════
#  5. LE RESTE DU FORMULAIRE NE BOUGE PAS
# ══════════════════════════════════════════════════════════════════════════

def test_la_grille_est_choisie_sur_ses_INTITULES_pas_sur_sa_forme():
    """POURQUOI CETTE RÈGLE EXISTE, ET POURQUOI ELLE FABRIQUE SON DOCUMENT.

    Le DC1 d'aujourd'hui ne contient QU'UNE table à trois colonnes : le cadre
    E. Une mutation qui supprime l'exigence des intitulés — « prends la
    première grille à trois colonnes » — ne change donc rien sur ce
    modèle-là, et survit sans que le garde-fou soit pour autant inutile : il
    protège contre la prochaine version du formulaire, où une deuxième grille
    apparaîtrait.

    ON FABRIQUE DONC LE CAS : deux tables à trois colonnes, un leurre d'abord.
    C'est ce qui rend le garde-fou MESURABLE au lieu d'être une intention."""
    from docx import Document
    doc = Document()
    leurre = doc.add_table(rows=2, cols=3)
    for i, t in enumerate(("Phase", "Début", "Fin")):
        leurre.rows[0].cells[i].text = t
    vraie = doc.add_table(rows=2, cols=3)
    for i, t in enumerate(("N° du Lot",
                           "Nom commercial et dénomination sociale, adresse",
                           "Prestations exécutées par les membres")):
        vraie.rows[0].cells[i].text = t

    rap = ao_formulaires.remplir_cadre_e(
        doc, [{"nom": "EPOC", "role": "Mandataire", "prestations": "Conception"}])
    assert rap["ecrits"] == 1
    assert vraie.rows[1].cells[1].text == "EPOC"
    assert [c.text for c in leurre.rows[1].cells] == ["", "", ""], \
        "les membres ont été écrits dans le premier tableau venu"


def test_une_grille_introuvable_est_DITE_et_rien_n_est_ecrit():
    """UN DOCUMENT SANS CADRE E N'EST PAS UN DOCUMENT SANS MEMBRES. Les deux
    motifs ne se soignent pas pareil : l'un demande une note, l'autre un
    modèle. Les confondre ferait chercher la note quand c'est le formulaire
    qui a changé."""
    from docx import Document
    rap = ao_formulaires.remplir_cadre_e(
        Document(), [{"nom": "EPOC", "role": "", "prestations": "x"}])
    assert rap["motif"] == "grille_introuvable"
    assert rap["ecrits"] == 0


def test_l_entete_de_la_grille_est_INTACTE():
    """LES INTITULÉS DE COLONNES SONT CEUX DU MINISTÈRE. Les réécrire — même
    à l'identique — ferait perdre leur mise en forme, et c'est par eux que la
    grille est retrouvée au tour suivant."""
    avant = [c.text for c in _grille(_modele()).rows[0].cells]
    doc, _ = _produire(ao_dc.membres_du_groupement(NOTE))
    assert [c.text for c in _grille(doc).rows[0].cells] == avant


def test_AUCUNE_autre_table_du_DC1_n_est_touchee():
    """L'INVARIANT DE CE MODULE : le formulaire produit ne diffère du modèle
    que là où une valeur est posée. Le cadre E est la SEULE grille que ce
    module écrit ; les autres — dont les cadres de déclaration sur l'honneur
    et de signature — doivent ressortir au caractère près."""
    mod = _modele()
    doc, _ = _produire(ao_dc.membres_du_groupement(NOTE))
    saut = _rang_grille(mod)
    assert saut is not None, "la grille du cadre E n'existe plus dans le modèle"
    assert len(doc.tables) == len(mod.tables), \
        "le remplissage a ajouté ou supprimé une table entière"
    ecarts = []
    for i, (a, b) in enumerate(zip(mod.tables, doc.tables)):
        if i == saut:
            continue
        ta = [c.text for r in a.rows for c in r.cells]
        tb = [c.text for r in b.rows for c in r.cells]
        if ta != tb:
            ecarts.append(i)
    assert not ecarts, "table(s) modifiée(s) hors du cadre E : %r" % ecarts


def test_les_membres_ne_changent_RIEN_au_placement_des_valeurs():
    """LES DEUX MÉCANISMES SONT INDÉPENDANTS, ET C'EST LA VRAIE PROPRIÉTÉ.

    CE QUE CETTE RÈGLE REMPLACE. Le module portait un commentaire affirmant
    que le cadre E DOIT être rempli avant le relevé des paragraphes, « parce
    qu'ajouter des rangées déplace les paragraphes qui suivent ». Une mutation
    qui intervertit les deux lignes a SURVÉCU — et la mesure a montré que
    l'affirmation était fausse : les deux ordres rendent le même document au
    caractère près. Écrire une règle pour « tenir l'ordre » aurait donné une
    règle verte qui ne mesure rien, exactement le défaut que ce dossier
    poursuit.

    CE QUI COMPTE, LUI, SE MESURE : aucune valeur d'ancre ne doit atterrir
    dans la grille des membres, et la grille ne doit pas faire manquer une
    ancre. Le jour où une ancre ou une zone interdite tomberait dans le cadre
    E, c'est ICI que cela se verrait."""
    # AUCUNE DE CES VALEURS N'EST UN NOM DE MEMBRE, et c'est délibéré : avec
    # « CONSEILPREV » comme dénomination du candidat — ce qu'elle est dans la
    # vraie vie, et ce qu'elle était ici —, une valeur d'ancre égarée dans le
    # cadre E se confondait avec le cotraitant du même nom, et la règle ne
    # voyait rien. Une mutation qui ajoute une ancre pointant DANS la grille a
    # survécu pour cette seule raison.
    valeurs = {"acheteur": "Ville de X", "objet_consultation": "Maîtrise d'œuvre",
               "candidat": "ZZZ Cabinet témoin", "adresse": "1 rue A, 75000 Paris",
               "courriel": "temoin@x.fr", "telephone": "01 02 03 04 05",
               "siret": "73282932000074"}
    _o, sans = ao_formulaires.remplir_document("dc1", valeurs, membres=[])
    octets, avec = ao_formulaires.remplir_document(
        "dc1", valeurs, membres=ao_dc.membres_du_groupement(NOTE))
    assert [x["rubrique"] for x in avec["places"]] == \
        [x["rubrique"] for x in sans["places"]]
    assert avec["non_places"] == sans["non_places"]
    assert avec["ignores"] == sans["ignores"]

    # ET RÉCIPROQUEMENT : la grille ne contient QUE ce que la note a donné.
    from docx import Document
    t = _grille(Document(io.BytesIO(octets)))
    membres = ao_dc.membres_du_groupement(NOTE)
    attendu = {""} | {m["nom"] for m in membres} | {m["prestations"] for m in membres}
    trouve = {c.text.strip() for r in t.rows[1:] for c in r.cells}
    intrus = trouve - attendu
    assert not intrus, "une valeur d'ancre a atterri dans le cadre E : %r" % intrus


def test_sans_membre_le_document_est_celui_du_modele():
    """LE TÉMOIN NÉGATIF DU DOCUMENT. Sans lui, une règle qui constate quatre
    lignes écrites passerait aussi sur un module qui ajoute toujours des
    rangées, membres ou pas."""
    mod = _grille(_modele())
    doc, rap = _produire([])
    t = _grille(doc)
    assert rap["cadre_e"]["motif"] == "aucun_membre"
    assert rap["cadre_e"]["lignes_ajoutees"] == 0
    assert len(t.rows) == len(mod.rows)
    assert [c.text for r in t.rows for c in r.cells] == \
        [c.text for r in mod.rows for c in r.cells]


def test_les_trois_autres_formulaires_n_ont_pas_de_cadre_E():
    """DC2, DC4 ET ATTRI1 NE DEMANDENT PAS LA RÉPARTITION. Leur passer des
    membres ne doit RIEN écrire — et le rapport doit le dire « sans objet »,
    pas « aucun membre » : les deux motifs ne se soignent pas pareil."""
    membres = ao_dc.membres_du_groupement(NOTE)
    for cle in sorted(ao_formulaires.MODELES):
        if cle == "dc1":
            continue
        _o, rap = ao_formulaires.remplir_document(cle, {}, membres=membres)
        assert rap["cadre_e"]["motif"] == "sans_objet", cle
        assert rap["cadre_e"]["ecrits"] == 0, cle


# ══════════════════════════════════════════════════════════════════════════
#  6. LA ROUTE, L'ARCHIVE ET LA CARTE LE DISENT
# ══════════════════════════════════════════════════════════════════════════

def _charge(piece="dc1"):
    return {"piece": piece, "fiche": FICHE, "format": "docx",
            "groupement": True, "analyse": _analyse_avec_note()}


def test_la_route_piece_rend_le_cadre_E_dans_son_en_tete(marche):
    """LA CARTE N'A PAS D'AUTRE SOURCE. Si la route ne le dit pas, la page ne
    peut pas le montrer, et cinq cotraitants écrits ressemblent à un DC1
    ordinaire."""
    r = marche.post("/api/datacenter/marche/piece", json=_charge(),
                    headers=ORIGINE)
    assert r.status_code == 200
    ce = json.loads(r.headers["X-Piece"])["cadre_e"]
    assert ce["ecrits"] == 4
    assert ce["a_confirmer"] == AVEC_RESERVE


def test_le_document_servi_par_la_route_porte_bien_les_membres(marche):
    """ON OUVRE CE QUE LA ROUTE A SERVI, pas ce que le module a rendu à côté.
    C'est la seule mesure qui couvre le branchement `membres=` de la route."""
    from docx import Document
    r = marche.post("/api/datacenter/marche/piece", json=_charge(),
                    headers=ORIGINE)
    t = _grille(Document(io.BytesIO(r.data)))
    assert [x.cells[1].text.strip() for x in t.rows[1:] if x.cells[1].text.strip()] \
        == NOMS_ATTENDUS


def test_le_bordereau_de_l_archive_NOMME_le_cadre_E(marche):
    """CELUI QUI REÇOIT L'ARCHIVE NE LIT PAS LES CARTES. Le bordereau est tout
    ce qu'il a : un DC1 sorti avec ses cotraitants dedans doit s'y voir, et les
    réserves encore ouvertes aussi."""
    r = marche.post("/api/datacenter/marche/dossier.zip", json=_charge(),
                    headers=ORIGINE)
    assert r.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r.data)) as z:
        bord = z.read("BORDEREAU.txt").decode("utf-8")
        t = _grille(__import__("docx").Document(
            io.BytesIO(z.read("dc1-projet-non-signe.docx"))))
    ligne = [l for l in bord.splitlines() if "dc1-projet-non-signe" in l]
    assert ligne, "le DC1 n'est pas au bordereau"
    assert "cadre E : 4 membre(s)" in ligne[0], ligne[0]
    assert "3 à confirmer" not in ligne[0] and "2 à confirmer" in ligne[0]
    assert [x.cells[1].text.strip() for x in t.rows[1:]
            if x.cells[1].text.strip()] == NOMS_ATTENDUS


def test_le_bordereau_se_TAIT_quand_aucun_membre_n_est_lu(marche):
    """LA MENTION EST UN FAIT, PAS UNE RUBRIQUE FIXE. « cadre E : 0 membre »
    sur tous les dossiers sans groupement serait du bruit, et le jour où elle
    compterait vraiment on ne la verrait plus."""
    charge = _charge()
    charge["analyse"] = ao_dc.analyser([])
    r = marche.post("/api/datacenter/marche/dossier.zip", json=charge,
                    headers=ORIGINE)
    with zipfile.ZipFile(io.BytesIO(r.data)) as z:
        bord = z.read("BORDEREAU.txt").decode("utf-8")
    assert "cadre E" not in bord


def test_la_carte_montre_les_membres_et_les_reserves():
    """LA FONCTION DE LA PAGE EST EXÉCUTÉE, pas relue. Une règle qui cherche
    « cadre E » dans la source passerait sur un bouton mort."""
    src = _js_source("aoCadreE")
    js = """
    function esc(s){ return String(s).replace(/&/g,"&amp;").replace(/</g,"&lt;"); }
    %s
    var vide = aoCadreE(null) + aoCadreE({ecrits: 0, a_confirmer: ["X"]});
    var plein = aoCadreE({ecrits: 4, lignes_ajoutees: 1,
                          a_confirmer: %s,
                          a_completer: ["N° du lot par membre"]});
    console.log(JSON.stringify({vide: vide, plein: plein}));
    """ % (src, json.dumps(AVEC_RESERVE, ensure_ascii=False))
    import subprocess
    out = subprocess.run(["/opt/node22/bin/node", "-e", js],
                         capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stderr
    d = json.loads(out.stdout)
    assert d["vide"] == "", "la carte parle d'un cadre E qui n'a rien reçu"
    assert "4 membre(s)" in d["plein"]
    for n in AVEC_RESERVE:
        assert n in d["plein"], n
    assert "lot" in d["plein"]


def test_la_carte_appelle_aoCadreE_la_ou_elle_montre_la_production():
    """LA FONCTION POSÉE DOIT ÊTRE LUE. Une fonction juste que personne
    n'appelle est du code mort qui rend une règle verte — le défaut nommé
    plusieurs fois dans ce dossier."""
    src = io.open(os.path.join(ICI, "ingenierie-dc.js"), encoding="utf-8").read()
    corps = _js_source("aoLotEtatCarte")
    assert "aoCadreE(" in corps, \
        "aoLotEtatCarte ne montre pas le cadre E"
    assert re.search(r"groupement:\s*\(d && d\.cadre_e\)", src), \
        "la production ne garde pas le cadre E rendu par la route"
